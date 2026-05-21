"""Генератор Excel-отчётов в формате ФСТЭК."""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO
from typing import List, Dict, Any
from app.models.schemas import ExcelRow


# Заголовки столбцов в формате ФСТЭК (Таблица 11)
COLUMNS = [
    "Идентификатор УБИ",
    "Наименование УБИ",
    "Уровень нарушителя (Внутренний)",
    "Уровень нарушителя (Внешний)",
    "Объект воздействия",
    "Способы реализации",
    "Негативные последствия",
    "Тактика",
    "Техника",
    "Примечания"
]


def create_excel_header_style():
    """Создать стиль для заголовков."""
    return Font(bold=True, size=12, name="Times New Roman")


def create_excel_cell_style():
    """Создать стиль для ячеек."""
    return Font(size=11, name="Times New Roman")


def create_border():
    """Создать границы для ячеек."""
    thin = Side(style="thin")
    return Border(
        left=thin, right=thin, top=thin, bottom=thin
    )


def format_list(items: List[Any]) -> str:
    """Преобразовать список в строку через запятую."""
    if not items:
        return ""
    return ", ".join(str(item) for item in items)


def generate_excel_report(threats_data: List[Dict[str, Any]]) -> BytesIO:
    """
    Сгенерировать Excel-файл с перечнем угроз.
    
    Args:
        threats_data: Список словарей с данными угроз
    
    Returns:
        BytesIO объект с Excel-файлом
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Перечень возможных (вероятных) УБИ для Системы"
    
    # Настройка ширины столбцов
    column_widths = [15, 45, 20, 20, 20, 25, 30, 15, 20, 30]
    for i, width in enumerate(column_widths, start=1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = width
    
    # Заголовки
    header_font = create_excel_header_style()
    cell_font = create_excel_cell_style()
    border = create_border()
    
    # Fill заголовков (светло-серый)
    header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    
    for col_num, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        cell.fill = header_fill
    
    # Данные
    for row_num, threat in enumerate(threats_data, start=2):
        # Идентификатор УБИ
        cell = ws.cell(row=row_num, column=1, value=threat.get("ubi_id", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border
        
        # Наименование УБИ
        cell = ws.cell(row=row_num, column=2, value=threat.get("ubi_name", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = border
        
        # Уровень нарушителя (Внутренний)
        cell = ws.cell(row=row_num, column=3, value=threat.get("violator_internal", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        
        # Уровень нарушителя (Внешний)
        cell = ws.cell(row=row_num, column=4, value=threat.get("violator_external", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        
        # Объект воздействия
        cell = ws.cell(row=row_num, column=5, value=threat.get("objects", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border
        
        # Способы реализации
        cell = ws.cell(row=row_num, column=6, value=threat.get("methods", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = border
        
        # Негативные последствия
        cell = ws.cell(row=row_num, column=7, value=threat.get("consequences", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = border
        
        # Тактика
        cell = ws.cell(row=row_num, column=8, value=threat.get("tactics", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border
        
        # Техника
        cell = ws.cell(row=row_num, column=9, value=threat.get("techniques", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border
        
        # Примечания
        cell = ws.cell(row=row_num, column=10, value=threat.get("notes", ""))
        cell.font = cell_font
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = border
    
    # Заморозить первую строку
    ws.freeze_panes = "A2"
    
    # Сохранить в BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output


def generate_excel_from_threats(threats: List[Any]) -> BytesIO:
    """
    Сгенерировать Excel из SQLAlchemy моделей Threat.
    
    Args:
        threats: Список моделей Threat
    
    Returns:
        BytesIO объект с Excel-файлом
    """
    threats_data = []
    
    for threat in threats:
        # Формируем строку данных
        data = {
            "ubi_id": threat.id,
            "ubi_name": threat.name,
            "violator_internal": format_list([v.level for v in threat.violator_int]),
            "violator_external": format_list([v.level for v in threat.violator_ext]),
            "objects": format_list([obj.code for obj in threat.objects]),
            "methods": format_list([m.code for m in threat.methods]),
            "consequences": format_list([c.code for c in threat.consequences]),
            "tactics": format_list([t.code for t in threat.tactics]),
            "techniques": format_list([tech.code for tech in threat.techniques]),
            "notes": threat.exclusion_note or ""
        }
        threats_data.append(data)
    
    return generate_excel_report(threats_data)


class ExcelExporter:
    """Экспортёр отчётов в Excel."""
    
    def __init__(self):
        self.columns = COLUMNS
    
    def export(self, threats_data: List[Dict[str, Any]]) -> BytesIO:
        """Экспортировать данные в Excel."""
        return generate_excel_report(threats_data)
    
    def export_from_models(self, threats: List[Any]) -> BytesIO:
        """Экспортировать SQLAlchemy модели в Excel."""
        return generate_excel_from_threats(threats)
